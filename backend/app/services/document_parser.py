import os
import io
import email
from email import policy
from email.parser import BytesParser
from typing import Tuple

class DocumentParserService:
    @staticmethod
    def parse_document(file_bytes: bytes, filename: str) -> Tuple[str, str]:
        """
        Extracts plain text from PDF, DOCX, TXT, or EML files.
        Returns (extracted_text, file_type).
        """
        ext = os.path.splitext(filename)[1].lower().strip(".")
        
        if ext == "pdf":
            text = DocumentParserService._extract_pdf(file_bytes)
            return text, "PDF"
        elif ext in ["docx", "doc"]:
            text = DocumentParserService._extract_docx(file_bytes)
            return text, "DOCX"
        elif ext in ["eml", "msg"]:
            text = DocumentParserService._extract_eml(file_bytes)
            return text, "EML"
        elif ext in ["txt", "log", "csv", "json"]:
            text = DocumentParserService._extract_txt(file_bytes)
            return text, "TXT"
        else:
            # Attempt plain text decoding as fallback
            text = DocumentParserService._extract_txt(file_bytes)
            return text, ext.upper() if ext else "UNKNOWN"

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text.strip())
            return "\n\n".join(pages_text).strip()
        except Exception as e:
            return f"Error extracting PDF text: {str(e)}"

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))
            
            all_text = "\n".join(paragraphs)
            if table_text:
                all_text += "\n\n--- Tables in Document ---\n" + "\n".join(table_text)
            return all_text.strip()
        except Exception as e:
            return f"Error extracting DOCX text: {str(e)}"

    @staticmethod
    def _extract_eml(file_bytes: bytes) -> str:
        try:
            msg = BytesParser(policy=policy.default).parsebytes(file_bytes)
            subject = msg.get("subject", "No Subject")
            sender = msg.get("from", "Unknown Sender")
            recipient = msg.get("to", "Unknown Recipient")
            date = msg.get("date", "")
            
            body = ""
            if msg.is_multipart():
                for part in msg.walk():
                    content_type = part.get_content_type()
                    content_disposition = str(part.get("Content-Disposition"))
                    if content_type == "text/plain" and "attachment" not in content_disposition:
                        body += part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", errors="replace") + "\n"
            else:
                body = msg.get_payload(decode=True).decode(msg.get_content_charset() or "utf-8", errors="replace")
                
            return f"From: {sender}\nTo: {recipient}\nDate: {date}\nSubject: {subject}\n\n{body.strip()}"
        except Exception as e:
            return f"Error extracting EML text: {str(e)}"

    @staticmethod
    def _extract_txt(file_bytes: bytes) -> str:
        for enc in ["utf-8", "utf-16", "latin-1", "cp1252"]:
            try:
                return file_bytes.decode(enc).strip()
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace").strip()

document_parser = DocumentParserService()
